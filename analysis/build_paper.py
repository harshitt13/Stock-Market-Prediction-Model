"""Concatenate the paper sections in docs/paper/ into one document.

Outputs, both under docs/paper/build/:

  paper_full.md    the ten section files in order, headings renumbered 1 to 9
                   with subsections 1.1, 1.2, ..., the section 2 reference
                   list moved to the end and deduplicated, every [VERIFY] and
                   [unsourced] marker lifted out of the prose into Appendix A
                   with the section it came from, and the author notes that
                   are not part of the paper (the section 2 verification log,
                   the section 3 venue note, the per-section sourcing
                   sentences that refer to the markers) removed and listed in
                   Appendix B.
  paper_full.docx  the same text through pandoc (pypandoc_binary), with a
                   title block and abstract on the first page, a page break
                   before section 1, the figures from docs/figures/ inserted
                   after the first paragraph that cites each one by filename,
                   captions taken from README section 4, and a centred page
                   number in the footer (python-docx).

The build fails loudly if a section file is missing or empty, if a section's
heading does not reach the output, if a cross-reference of the form
section-sign X.Y or "Table N" points at nothing, or if a referenced figure is
not on disk. It prints, but does not fail on, citations with no entry in the
reference list and entries that nothing cites.

Usage:
    python analysis/build_paper.py            # both outputs
    python analysis/build_paper.py --no-docx  # markdown only
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
PAPER = REPO / "docs" / "paper"
BUILD = PAPER / "build"
FIGURES = REPO / "docs" / "figures"
README = REPO / "README.md"

SECTIONS = [
    "00_abstract.md",
    "01_introduction.md",
    "02_related_work.md",
    "03_contributions.md",
    "04_artifacts.md",
    "05_method.md",
    "06_results.md",
    "07_discussion.md",
    "08_limitations.md",
    "09_reproducibility.md",
]
TITLE = (
    "Protocol Artifacts in Deep Learning for Stock Return Prediction: "
    "A Reproducible Null Result Across 30 Equities"
)
AUTHOR = "[Author names and affiliations]"

UNSOURCED = "**[unsourced]**"
# A sentence about the marking convention itself ("... is marked **[unsourced]**"),
# which has nothing to say once the markers are gone.
META_RE = re.compile(r"marked(?:\s+it)?\s+(?:\*\*\[unsourced\]\*\*|`\[VERIFY\]`)")
# A sentence ends at . ! or ?, optionally closed by ** or ), followed by
# whitespace and something that opens a sentence. Only paragraphs that carry a
# marker are split with this, so "et al. (2020)" in the survey is never touched.
SENT_END = re.compile(r"(?:(?<=[.!?])|(?<=[.!?]\*\*)|(?<=[.!?]\)))\s+(?=[A-Z*(§`\[])")
VERIFY_RE = re.compile(r"\s*`\[VERIFY(?::\s*(?P<note>[^\]]*))?\]`", re.S)
FIG_RE = re.compile(r"docs/figures/([A-Za-z0-9_]+\.png)")
CITE_RE = re.compile(
    r"([A-Z][\w'’\-]+)"                       # first surname
    r"(?:,? (?:and|&) [A-Z][\w'’\-]+"          # A and B
    r"|, [A-Z][\w'’\-]+,? (?:and|&) [A-Z][\w'’\-]+"  # A, B and C
    r"| et al\.)?"
    r"(?:'s|’s)? \((\d{4}[a-z]?)\)"
)
# Parenthetical citations: "(Sezer et al., 2020; Gu, Kelly and Xiu, 2020)".
PAREN_CITE_RE = re.compile(
    r"(?<![\w-])"                           # not the tail of "AAPL, 2010-03-16"
    r"((?=[A-Z][a-z])[A-Z][\w'’\-]+(?: [A-Z][\w'’\-]+)?)"  # a surname, not a ticker
    r"(?:,? (?:and|&) [A-Z][\w'’\-]+(?: [A-Z][\w'’\-]+)?"
    r"|, [A-Z][\w'’\-]+,? (?:and|&) [A-Z][\w'’\-]+"
    r"| et al\.)?"
    r", (\d{4}[a-z]?)(?=[;)])"                # the year closes the citation
)
XREF_RE = re.compile(r"(?<!README )§(\d+(?:\.\d+)*)")  # "README §3.5" points at the README
TABLE_REF_RE = re.compile(r"\bTable (\d+)")


def fail(msg: str) -> None:
    sys.exit(f"build_paper: {msg}")


def read_sections() -> dict[str, str]:
    out = {}
    for name in SECTIONS:
        path = PAPER / name
        if not path.is_file():
            fail(f"section file missing: {path}")
        text = path.read_text(encoding="utf-8").strip()
        if not text:
            fail(f"section file empty: {path}")
        out[name] = text
    return out


def paragraphs(text: str) -> list[str]:
    return re.split(r"\n\s*\n", text.strip())


def normalise_ws(s: str) -> str:
    return " ".join(s.split())


# --------------------------------------------------------------------------
# Section 2: detach the reference list and drop the verification log.
def split_related_work(text: str) -> tuple[str, list[str]]:
    if "## References" not in text:
        fail("02_related_work.md has no '## References' heading")
    body, rest = text.split("## References", 1)
    if "## Verification log" in rest:
        refs_block, _log = rest.split("## Verification log", 1)
    else:
        refs_block = rest
    refs_block = refs_block.strip().rstrip("-").strip()
    entries = re.findall(r"^- (.+?)(?=^- |\Z)", refs_block + "\n", flags=re.S | re.M)
    if not entries:
        fail("no reference entries parsed from 02_related_work.md")
    return body.rstrip().rstrip("-").rstrip(), [normalise_ws(e) for e in entries]


def dedupe_references(entries: list[str]) -> tuple[list[str], list[str]]:
    seen: dict[str, str] = {}
    dropped = []
    for e in entries:
        key = re.sub(r"[^a-z0-9]", "", e.lower())
        if key in seen:
            dropped.append(e)
        else:
            seen[key] = e
    kept = sorted(seen.values(), key=lambda s: s.lower())
    return kept, dropped


def reference_keys(entries: list[str]) -> set[tuple[str, str]]:
    keys = set()
    for e in entries:
        m = re.match(r"([A-Za-zÀ-ÿ'’\- ]+?),.*?\((\d{4}[a-z]?)\)", e)
        if m:
            keys.add((m.group(1).strip().split()[-1], m.group(2)))
    return keys


# --------------------------------------------------------------------------
# Headings.
def renumber(text: str, number: int | None) -> tuple[str, list[str]]:
    """Rewrite headings so the level-1 heading carries `number` and level-2
    headings carry number.k in order. Level-3 headings and the abstract stay
    unnumbered. Returns the text and the list of heading labels produced."""
    labels: list[str] = []
    k = 0
    out = []
    for line in text.split("\n"):
        m = re.match(r"^(#{1,3})\s+(?:\d+(?:\.\d+)*\.?\s+)?(.*)$", line)
        if not m:
            out.append(line)
            continue
        level, title = len(m.group(1)), m.group(2).strip()
        if level == 1 and number is not None:
            out.append(f"# {number}. {title}")
            labels.append(str(number))
        elif level == 2 and number is not None:
            k += 1
            out.append(f"## {number}.{k} {title}")
            labels.append(f"{number}.{k}")
        else:
            out.append(f"{'#' * level} {title}")
    return "\n".join(out), labels


# --------------------------------------------------------------------------
# Author notes and markers.
class Ledger:
    def __init__(self) -> None:
        self.markers: list[tuple[str, str, str]] = []  # (section, kind, text)
        self.notes: list[tuple[str, str]] = []         # (section, text)


def current_section(para: str, state: dict) -> str:
    m = re.match(r"^(#{1,2})\s+(\d+(?:\.\d+)?)\.?\s", para)
    if m:
        state["section"] = m.group(2)
    return state.get("section", "front matter")


def strip_notes_and_markers(text: str, ledger: Ledger, state: dict) -> str:
    out = []
    for para in paragraphs(text):
        sec = current_section(para, state)
        # Whole-paragraph author notes: "*Venue note (not part of the paper) ...*"
        if re.match(r"^\*[^*\n]*\(not part of the paper\)", para):
            ledger.notes.append((sec, normalise_ws(para.strip("*"))))
            continue
        # Sentences that only describe the marking convention.
        if META_RE.search(para):
            kept = []
            for s in SENT_END.split(para):
                if META_RE.search(s):
                    ledger.notes.append((sec, normalise_ws(s)))
                else:
                    kept.append(s)
            para = " ".join(kept).strip()
            if not para:
                continue
        # [VERIFY] brackets.
        def _verify(m: re.Match) -> str:
            note = m.group("note")
            if note:
                ledger.markers.append((sec, "[VERIFY]", normalise_ws(note)))
            else:
                line = para[: m.start()].rsplit("\n", 1)[-1] + para[m.start():].split("\n", 1)[0]
                first_cell = line.strip().strip("|").split("|")[0].strip() if line.lstrip().startswith("|") else ""
                ledger.markers.append((sec, "[VERIFY]", f"table row {first_cell!r}, target column"))
            return ""
        para = VERIFY_RE.sub(_verify, para)
        # [unsourced] flags: lift the marker, keep the sentence, capitalise it.
        while UNSOURCED + ":" in para:
            i = para.index(UNSOURCED + ":")
            after = para[i + len(UNSOURCED) + 1:].lstrip()
            m = SENT_END.search(after)
            sentence = after[: m.start()] if m else after
            ledger.markers.append((sec, "[unsourced]", normalise_ws(sentence)))
            rest = after[len(sentence):]
            head = para[:i]
            if head.rstrip().endswith((".", ":")) or not head.strip():
                sentence = sentence[0].upper() + sentence[1:]
            para = head + sentence + rest
        if UNSOURCED in para or "[VERIFY" in para:
            fail(f"a marker survived the build in section {sec}: {normalise_ws(para)[:120]!r}")
        out.append(para)
    return "\n\n".join(out)


# --------------------------------------------------------------------------
# Figures.
def readme_captions() -> dict[str, str]:
    caps = {}
    for line in README.read_text(encoding="utf-8").split("\n"):
        if not line.startswith("|") or ".png" not in line:
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 2:
            continue
        m = FIG_RE.search(cells[0])
        if not m:
            continue
        cap = cells[1]
        cap = re.sub(r"^\*\*`[^`]+`\*\*\s*[—-]\s*", "", cap)  # "**`name.png`** — "
        cap = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", cap)
        cap = cap.replace("**", "").replace("`", "").replace("\\|", "|")
        caps[m.group(1)] = normalise_ws(cap)
    return caps


def insert_figures(text: str, captions: dict[str, str], placed: list[str]) -> str:
    out = []
    for para in paragraphs(text):
        out.append(para)
        for name in FIG_RE.findall(para):
            if name in placed:
                continue
            if not (FIGURES / name).is_file():
                fail(f"figure referenced in the text is not on disk: docs/figures/{name}")
            placed.append(name)
            n = len(placed)
            cap = captions.get(name, "")
            if not cap:
                print(f"  warning: no README caption for {name}; using the filename")
            caption = f"Figure {n}. {name}" + (f": {cap}" if cap else "")
            out.append(f"![{caption}](../../figures/{name}){{width=6in}}")
    return "\n\n".join(out)


# --------------------------------------------------------------------------
# Checks.
def check_cross_references(text: str, labels: set[str]) -> None:
    bad = set()
    for m in XREF_RE.finditer(text):
        if m.group(1) in labels:
            continue
        # References to the README come in lists whose first item alone carries
        # the prefix: "(README §5.1, §12)", "`README.md` §3.1". Look at the clause.
        clause = re.split(r"[(.;]\s|\n\n", text[max(0, m.start() - 120):m.start()])[-1]
        if "README" not in clause:
            bad.add(m.group(1))
    bad = sorted(bad)
    if bad:
        fail(f"cross-references to headings that do not exist: {bad}")
    captions = set(re.findall(r"\*\*Table (\d+)\.\*\*", text))
    bad_t = sorted({t for t in TABLE_REF_RE.findall(text) if t not in captions})
    if bad_t:
        fail(f"references to tables without a '**Table N.**' caption: {bad_t}")


def report_citations(body: str, entries: list[str]) -> None:
    keys = reference_keys(entries)
    flat = normalise_ws(body)
    cited = set()
    for rx in (CITE_RE, PAREN_CITE_RE):
        for m in rx.finditer(flat):
            # "Diebold-Mariano (1995)" keys on Diebold; "Shafiei Hafshejani" on Hafshejani.
            surname = re.split(r"[-–]", m.group(1))[0].split()[-1]
            cited.add((surname, m.group(2)))
    missing = sorted(c for c in cited if c not in keys)
    unused = sorted(k for k in keys if k not in cited)
    print("  citations with no reference entry:", missing or "none")
    print("  reference entries nothing cites:", unused or "none")


# --------------------------------------------------------------------------
def build_markdown() -> tuple[str, dict]:
    files = read_sections()
    ledger = Ledger()
    state: dict = {}
    captions = readme_captions()
    placed: list[str] = []
    parts = []
    labels: set[str] = set()
    level1_titles = []

    body2, entries = split_related_work(files["02_related_work.md"])
    files["02_related_work.md"] = body2

    for number, name in enumerate(SECTIONS):
        text = files[name]
        text, labs = renumber(text, None if number == 0 else number)
        labels.update(labs)
        text = strip_notes_and_markers(text, ledger, state)
        text = insert_figures(text, captions, placed)
        level1_titles.append(text.split("\n", 1)[0])
        parts.append(text)

    refs, dropped = dedupe_references(entries)
    parts.append("# References\n\n" + "\n".join(f"- {e}" for e in refs))

    app_a = ["# Appendix A. Outstanding verification markers", "",
             "The following markers were lifted out of the text by "
             "`analysis/build_paper.py`. Each `[unsourced]` sentence remains in "
             "the section named, without its marker; each `[VERIFY]` note was "
             "removed from the section named.", "",
             "| # | Section | Marker | Text |", "|---|---|---|---|"]
    for i, (sec, kind, txt) in enumerate(ledger.markers, 1):
        app_a.append(f"| {i} | §{sec} | {kind} | {txt.replace('|', '\\|')} |")
    app_a += ["", f"Title page: {AUTHOR} is a placeholder."]
    parts.append("\n".join(app_a))

    app_b = ["# Appendix B. Notes removed from the text", "",
             "Author notes that are not part of the paper, removed by the build:", ""]
    app_b.append("- §2: the verification log that follows the reference list in "
                 "`docs/paper/02_related_work.md` (the checking record for every "
                 "characterisation in §2).")
    for sec, txt in ledger.notes:
        app_b.append(f"- §{sec}: {txt}")
    if dropped:
        app_b.append(f"- Duplicate reference entries removed: {len(dropped)}.")
    parts.append("\n".join(app_b))

    today = dt.date.today()
    front = ["---", f'title: "{TITLE}"', f'author: "{AUTHOR}"',
             f'date: "{today.day} {today:%B %Y}"', "---", ""]
    md = "\n".join(front) + "\n" + "\n\n".join(parts) + "\n"

    # Loud checks.
    for name, title in zip(SECTIONS, level1_titles):
        if title not in md:
            fail(f"heading of {name} did not reach the output: {title!r}")
    check_cross_references(md, labels)
    print(f"  sections: {len(SECTIONS)}, headings: {len(labels)}, figures placed: {placed}")
    print(f"  markers lifted: {len(ledger.markers)}, notes removed: {len(ledger.notes)}, "
          f"references: {len(refs)} (duplicates dropped: {len(dropped)})")
    report_citations("\n".join(parts[:-3]), refs)
    body_words = len("\n\n".join(parts[:-3]).split())
    info = {"words": len(md.split()), "body_words": body_words, "figures": len(placed),
            "table_rows": sum(1 for l in md.split("\n") if l.startswith("|") and not set(l) <= set("|-: ")),
            "refs": len(refs), "markers": len(ledger.markers)}
    return md, info


def build_docx(md: str, out: Path) -> None:
    import pypandoc  # pypandoc_binary ships the pandoc executable

    page_break = "```{=openxml}\n<w:p><w:r><w:br w:type=\"page\"/></w:r></w:p>\n```\n\n"
    md_docx = md.replace("\n# 1. ", "\n" + page_break + "# 1. ", 1)
    pypandoc.convert_text(
        md_docx, "docx", format="markdown", outputfile=str(out),
        extra_args=["--resource-path", str(BUILD), "--standalone"],
    )
    # Page numbers: a centred PAGE field in every section footer.
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(str(out))
    for section in doc.sections:
        p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        run._r.append(fld)
    doc.save(str(out))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-docx", action="store_true", help="write paper_full.md only")
    args = ap.parse_args()
    BUILD.mkdir(parents=True, exist_ok=True)
    print("building paper_full.md")
    md, info = build_markdown()
    (BUILD / "paper_full.md").write_text(md, encoding="utf-8")
    # Page estimate: single-spaced 12 pt Letter with 1-inch margins runs about
    # 500 words of prose a page; a table row takes a line of about 45 a page;
    # a figure at 6 in wide takes about half a page with its caption.
    pages = info["words"] / 500 + info["table_rows"] / 45 + info["figures"] * 0.5
    print(f"  words: {info['words']:,} in total, {info['body_words']:,} in sections 1 to 9 and the abstract  "
          f"table rows: {info['table_rows']}  figures: {info['figures']}  estimated pages: {pages:.0f}")
    if not args.no_docx:
        print("building paper_full.docx")
        build_docx(md, BUILD / "paper_full.docx")
        print(f"  wrote {BUILD / 'paper_full.docx'}")


if __name__ == "__main__":
    main()
